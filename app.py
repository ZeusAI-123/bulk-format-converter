"""
Bulk Format Converter — Streamlit App (Zeus Doc Chat UI)
---------------------------------------------------
Supports: Images, Audio, Documents (docx, doc, pdf, rtf, odt, txt, html, md, csv, xlsx)

Run:
    streamlit run bulk_converter_zeus.py

Install (core):
    pip install streamlit Pillow pydub python-docx pymupdf pypandoc openpyxl odfpy html2text markdown

Install (optional but recommended):
    # Audio codec support:
    #   macOS:   brew install ffmpeg
    #   Ubuntu:  sudo apt-get install ffmpeg
    #   Windows: https://ffmpeg.org/download.html

    # Broad document conversion (pandoc):
    #   macOS:   brew install pandoc
    #   Ubuntu:  sudo apt-get install pandoc
    #   Windows: https://pandoc.org/installing.html

    # Legacy .doc support (LibreOffice):
    #   macOS:   brew install --cask libreoffice
    #   Ubuntu:  sudo apt-get install libreoffice
    #   Windows: https://www.libreoffice.org/download/download/
    #
    # Streamlit Cloud deployment note:
    #   Add a packages.txt file to your repo containing:
    #       libreoffice
    #       pandoc
    #   Streamlit Cloud will install them automatically via apt-get.
"""

import io
import shutil
import subprocess
import zipfile
import tempfile
from pathlib import Path
from datetime import datetime

import streamlit as st


# ============================================================
# BACKEND — format definitions (UNCHANGED)
# ============================================================

IMAGE_FORMATS = {
    ".png", ".jpg", ".jpeg", ".bmp", ".gif",
    ".tiff", ".tif", ".webp", ".ico", ".ppm",
}

AUDIO_FORMATS = {
    ".wav", ".mp3", ".ogg", ".flac", ".aac", ".wma", ".m4a",
}

DOC_FORMATS = {
    ".docx", ".doc", ".dotx", ".dotm", ".docm",
    ".pdf",
    ".rtf", ".odt",
    ".txt", ".md", ".html", ".htm",
    ".csv", ".xlsx",
}

VALID_IMAGE_OUTPUTS = {"png", "jpg", "jpeg", "bmp", "gif", "tiff", "webp"}
VALID_AUDIO_OUTPUTS = {"mp3", "wav", "ogg", "flac", "aac"}
VALID_DOC_OUTPUTS   = {"txt", "docx", "pdf", "html", "md", "rtf", "odt", "csv", "xlsx"}

DEFAULT_QUALITY = {
    "jpg": 85, "jpeg": 85, "webp": 80, "png": None,
    "mp3": 192, "ogg": 192, "aac": 192, "wav": None, "flac": None,
}

PANDOC_INPUTS  = {".docx", ".doc", ".rtf", ".odt", ".html", ".htm", ".md", ".txt", ".pdf"}
PANDOC_OUTPUTS = {"docx", "pdf", "rtf", "odt", "html", "md", "txt"}


def detect_file_type(ext: str):
    ext = ext.lower()
    if ext in IMAGE_FORMATS: return "image"
    if ext in AUDIO_FORMATS: return "audio"
    if ext in DOC_FORMATS:   return "doc"
    return None


def get_valid_outputs(file_type: str) -> set:
    return {
        "image": VALID_IMAGE_OUTPUTS,
        "audio": VALID_AUDIO_OUTPUTS,
        "doc":   VALID_DOC_OUTPUTS,
    }.get(file_type, set())


# ============================================================
# DEPENDENCY CHECKS (UNCHANGED)
# ============================================================

def _check_libreoffice() -> tuple[bool, str]:
    soffice = shutil.which("soffice")
    if soffice:
        return True, soffice
    candidates = [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/snap/bin/libreoffice",
        "/usr/bin/soffice",
        "/usr/lib/libreoffice/program/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for path in candidates:
        if Path(path).exists():
            return True, path
    return False, (
        "LibreOffice not found. Install it and make sure `soffice` is on PATH.\n"
        "  macOS:   brew install --cask libreoffice\n"
        "  Ubuntu:  sudo apt-get install libreoffice\n"
        "  Windows: https://www.libreoffice.org/download/download/\n"
        "  Streamlit Cloud: add 'libreoffice' to packages.txt"
    )


def _check_pandoc() -> tuple[bool, str]:
    try:
        import pypandoc
        path = pypandoc.get_pandoc_path()
        return True, path
    except OSError:
        pass
    except ImportError:
        pass
    pandoc = shutil.which("pandoc")
    if pandoc:
        return True, pandoc
    return False, (
        "Pandoc not found. Install it:\n"
        "  macOS:   brew install pandoc\n"
        "  Ubuntu:  sudo apt-get install pandoc\n"
        "  Windows: https://pandoc.org/installing.html\n"
        "  Streamlit Cloud: add 'pandoc' to packages.txt"
    )


# ============================================================
# IMAGE CONVERSION (UNCHANGED)
# ============================================================

def convert_image(input_path: Path, output_path: Path, quality):
    from PIL import Image
    img = Image.open(input_path)
    out_ext = output_path.suffix.lower()
    if out_ext in (".jpg", ".jpeg") and img.mode in ("RGBA", "P", "LA"):
        img = img.convert("RGB")
    if out_ext == ".png" and img.mode == "P":
        img = img.convert("RGBA")
    save_kwargs = {}
    if quality and out_ext in (".jpg", ".jpeg", ".webp"):
        save_kwargs["quality"] = quality
    img.save(output_path, **save_kwargs)


# ============================================================
# AUDIO CONVERSION (UNCHANGED)
# ============================================================

def convert_audio(input_path: Path, output_path: Path, bitrate):
    from pydub import AudioSegment
    audio = AudioSegment.from_file(str(input_path))
    export_kwargs = {"format": output_path.suffix.lstrip(".")}
    if bitrate and output_path.suffix.lower() in (".mp3", ".ogg", ".aac"):
        export_kwargs["bitrate"] = f"{bitrate}k"
    audio.export(str(output_path), **export_kwargs)


# ============================================================
# DOCUMENT HELPERS (UNCHANGED)
# ============================================================

def _pdf_to_text(input_path: Path) -> str:
    import fitz
    doc = fitz.open(str(input_path))
    return "\n".join(page.get_text() for page in doc)


def _read_docx_text(input_path: Path) -> str:
    from docx import Document
    doc = Document(str(input_path))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx_as_csv(input_path: Path) -> str:
    import openpyxl, csv, io as _io
    wb = openpyxl.load_workbook(str(input_path), data_only=True)
    ws = wb.active
    buf = _io.StringIO()
    writer = csv.writer(buf)
    for row in ws.iter_rows(values_only=True):
        writer.writerow([("" if v is None else str(v)) for v in row])
    return buf.getvalue()


def _csv_to_xlsx(input_path: Path, output_path: Path):
    import openpyxl, csv
    wb = openpyxl.Workbook()
    ws = wb.active
    with open(str(input_path), newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            ws.append(row)
    wb.save(str(output_path))


def _try_pandoc(input_path: Path, output_path: Path):
    pandoc_ok, pandoc_msg = _check_pandoc()
    if not pandoc_ok:
        raise RuntimeError(pandoc_msg)
    import pypandoc
    out_fmt = output_path.suffix.lstrip(".")
    pypandoc.convert_file(
        str(input_path),
        to=out_fmt,
        outputfile=str(output_path),
        extra_args=["--standalone"] if out_fmt == "html" else [],
    )


def _doc_to_docx_via_libreoffice(input_path: Path, work_dir: Path) -> Path:
    lo_ok, lo_path = _check_libreoffice()
    if not lo_ok:
        raise RuntimeError(lo_path)
    lo_profile = work_dir / "lo_profile"
    lo_profile.mkdir(exist_ok=True)
    profile_url = lo_profile.as_uri()
    cmd = [
        lo_path,
        f"-env:UserInstallation={profile_url}",
        "--headless",
        "--norestore",
        "--nofirststartwizard",
        "--convert-to", "docx:MS Word 2007 XML",
        "--outdir", str(work_dir),
        str(input_path),
    ]
    import os
    safe_env = {k: v for k, v in os.environ.items()
                if k not in ("HOME", "USERPROFILE")}
    safe_env["HOME"] = str(work_dir)
    result = subprocess.run(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=120,
        env=safe_env,
    )
    if result.returncode != 0:
        stderr_text = result.stderr.decode(errors="replace").strip()
        raise RuntimeError(
            f"LibreOffice conversion failed (exit {result.returncode}).\n{stderr_text}"
        )
    expected = work_dir / f"{input_path.stem}.docx"
    if not expected.exists():
        matches = list(work_dir.glob("*.docx"))
        if not matches:
            raise RuntimeError(
                f"LibreOffice ran successfully but no .docx was found in {work_dir}."
            )
        expected = matches[0]
    return expected


def _check_mammoth() -> bool:
    try:
        import mammoth
        return True
    except ImportError:
        return False


def _check_markdownify() -> bool:
    try:
        import markdownify
        return True
    except ImportError:
        return False


def _check_reportlab() -> bool:
    try:
        from reportlab.pdfgen import canvas
        return True
    except ImportError:
        return False


def _check_antiword() -> tuple[bool, str]:
    path = shutil.which("antiword")
    if path:
        return True, path
    return False, ""


def _doc_extract_html_via_mammoth(input_path: Path) -> str:
    import mammoth
    with open(input_path, "rb") as f:
        result = mammoth.convert_to_html(f)
    return result.value


def _doc_extract_text_via_mammoth(input_path: Path) -> str:
    import mammoth
    with open(input_path, "rb") as f:
        result = mammoth.extract_raw_text(f)
    return result.value


def _doc_extract_text_via_antiword(input_path: Path) -> str:
    aw_ok, aw_path = _check_antiword()
    if not aw_ok:
        raise RuntimeError("antiword not found")
    result = subprocess.run(
        [aw_path, str(input_path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=60,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode(errors="replace").strip())
    return result.stdout.decode(errors="replace")


def _text_to_pdf_via_reportlab(text: str, output_path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as rl_canvas
    c = rl_canvas.Canvas(str(output_path), pagesize=A4)
    width, height = A4
    margin = 2 * cm
    line_height = 14
    x = margin
    y = height - margin
    for raw_line in text.splitlines():
        for chunk_start in range(0, max(len(raw_line), 1), 100):
            line = raw_line[chunk_start:chunk_start + 100]
            if y < margin:
                c.showPage()
                y = height - margin
            c.setFont("Helvetica", 10)
            c.drawString(x, y, line)
            y -= line_height
    c.save()


def _html_to_docx(html: str, output_path: Path):
    import re
    from docx import Document
    text = re.sub(r"<br\s*/?>", "\n", html, flags=re.IGNORECASE)
    text = re.sub(r"<p[^>]*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"</p>", "", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    for ent, char in [("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"),
                      ("&nbsp;", " "), ("&quot;", '"'), ("&#39;", "'")]:
        text = text.replace(ent, char)
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(output_path))


def _convert_doc_file(input_path: Path, output_path: Path):
    out_ext = output_path.suffix.lower()
    lo_ok, lo_path = _check_libreoffice()

    _lo_err = None
    if lo_ok:
        try:
            with tempfile.TemporaryDirectory(prefix="bulk_conv_lo_") as lo_tmp:
                lo_dir = Path(lo_tmp)
                if out_ext == ".pdf":
                    lo_profile = lo_dir / "lo_profile"
                    lo_profile.mkdir(exist_ok=True)
                    import os
                    safe_env = {k: v for k, v in os.environ.items()
                                if k not in ("HOME", "USERPROFILE")}
                    safe_env["HOME"] = str(lo_dir)
                    cmd = [
                        lo_path,
                        f"-env:UserInstallation={lo_profile.as_uri()}",
                        "--headless", "--norestore", "--nofirststartwizard",
                        "--convert-to", "pdf",
                        "--outdir", str(lo_dir),
                        str(input_path),
                    ]
                    result = subprocess.run(cmd, stdout=subprocess.PIPE,
                                            stderr=subprocess.PIPE, timeout=120,
                                            env=safe_env)
                    if result.returncode != 0:
                        raise RuntimeError(
                            f"LibreOffice PDF export failed (exit {result.returncode}).\n"
                            + result.stderr.decode(errors="replace").strip()
                        )
                    pdf_out = lo_dir / f"{input_path.stem}.pdf"
                    if not pdf_out.exists():
                        matches = list(lo_dir.glob("*.pdf"))
                        if not matches:
                            raise RuntimeError("LibreOffice ran but produced no PDF.")
                        pdf_out = matches[0]
                    shutil.copy2(pdf_out, output_path)
                    return
                docx_path = _doc_to_docx_via_libreoffice(input_path, lo_dir)
                if out_ext == ".docx":
                    shutil.copy2(docx_path, output_path)
                    return
                convert_doc(docx_path, output_path)
                return
        except Exception as e:
            _lo_err = str(e)

    if _check_mammoth():
        try:
            if out_ext == ".txt":
                text = _doc_extract_text_via_mammoth(input_path)
                output_path.write_text(text, encoding="utf-8")
                return
            if out_ext == ".html":
                html = _doc_extract_html_via_mammoth(input_path)
                output_path.write_text(html, encoding="utf-8")
                return
            if out_ext == ".md":
                html = _doc_extract_html_via_mammoth(input_path)
                if _check_markdownify():
                    import markdownify
                    md = markdownify.markdownify(html, heading_style="ATX")
                else:
                    import re
                    md = re.sub(r"<[^>]+>", "", html)
                output_path.write_text(md, encoding="utf-8")
                return
            if out_ext == ".docx":
                html = _doc_extract_html_via_mammoth(input_path)
                _html_to_docx(html, output_path)
                return
            if out_ext == ".pdf":
                if _check_reportlab():
                    text = _doc_extract_text_via_mammoth(input_path)
                    _text_to_pdf_via_reportlab(text, output_path)
                    return
        except Exception as mammoth_err:
            _mammoth_err = str(mammoth_err)
    else:
        _mammoth_err = "mammoth not installed (pip install mammoth)"

    aw_ok, _ = _check_antiword()
    if aw_ok and out_ext == ".txt":
        text = _doc_extract_text_via_antiword(input_path)
        output_path.write_text(text, encoding="utf-8")
        return

    lo_install = (
        "  macOS:   brew install --cask libreoffice\n"
        "  Ubuntu:  sudo apt-get install libreoffice\n"
        "  Windows: https://www.libreoffice.org/download/download/\n"
        "  Streamlit Cloud: add 'libreoffice' to packages.txt"
    )
    if out_ext in (".rtf", ".odt"):
        raise RuntimeError(
            f".doc → {out_ext} requires LibreOffice (no pure-Python fallback exists).\n"
            + lo_install
        )
    if out_ext == ".pdf" and not _check_reportlab():
        raise RuntimeError(
            f".doc → .pdf: LibreOffice not found and reportlab is not installed.\n"
            f"  Fix A (best quality): install LibreOffice\n{lo_install}\n"
            f"  Fix B (plain-text PDF): pip install reportlab"
        )
    raise RuntimeError(
        f".doc → {out_ext}: no conversion tool available.\n"
        f"  Best fix:    install LibreOffice (full fidelity)\n{lo_install}\n"
        f"  Python fix:  pip install mammoth markdownify reportlab\n"
    )


def convert_doc(input_path: Path, output_path: Path):
    in_ext  = input_path.suffix.lower()
    out_ext = output_path.suffix.lower()
    out_fmt = out_ext.lstrip(".")

    LIBREOFFICE_ONLY_INPUTS = {".doc", ".dotx", ".dotm", ".docm"}
    if in_ext in LIBREOFFICE_ONLY_INPUTS:
        _convert_doc_file(input_path, output_path)
        return

    if in_ext == ".pdf":
        text = _pdf_to_text(input_path)
        if out_ext == ".txt":
            output_path.write_text(text, encoding="utf-8")
        elif out_ext in (".docx",):
            from docx import Document
            doc = Document()
            for line in text.split("\n"):
                doc.add_paragraph(line)
            doc.save(str(output_path))
        elif out_ext == ".html":
            html = (
                "<html><body>"
                + "".join(f"<p>{l}</p>" for l in text.split("\n") if l.strip())
                + "</body></html>"
            )
            output_path.write_text(html, encoding="utf-8")
        elif out_ext in (".md", ".rtf", ".odt"):
            _try_pandoc(input_path, output_path)
        else:
            raise ValueError(f"PDF → {out_ext} not supported")
        return

    if in_ext == ".csv" and out_ext == ".xlsx":
        _csv_to_xlsx(input_path, output_path)
        return

    if in_ext == ".xlsx" and out_ext == ".csv":
        output_path.write_text(_read_xlsx_as_csv(input_path), encoding="utf-8")
        return

    if in_ext == ".txt" and out_ext == ".docx":
        from docx import Document
        text = input_path.read_text(encoding="utf-8")
        doc  = Document()
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(str(output_path))
        return

    if in_ext == ".docx" and out_ext == ".txt":
        output_path.write_text(_read_docx_text(input_path), encoding="utf-8")
        return

    if in_ext in PANDOC_INPUTS and out_fmt in PANDOC_OUTPUTS:
        _try_pandoc(input_path, output_path)
        return

    raise ValueError(f"Unsupported document conversion: {in_ext} → {out_ext}")


# ============================================================
# RESULT CONTAINER (UNCHANGED)
# ============================================================

class ConversionResult:
    def __init__(self, source, target, status, error=""):
        self.source = source
        self.target = target
        self.status = status
        self.error  = error


# ============================================================
# ZEUS DOC CHAT UI — GLOBAL CSS
# ============================================================

st.set_page_config(
    page_title="BulkConvert",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded"
)

def inject_global_css():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&display=swap');

    :root {
        --bg-primary:    #071018;
        --bg-secondary:  #0f172a;
        --card-bg:       #111827;
        --text-primary:  #f9fafb;
        --text-secondary:#d1d5db;
        --text-muted:    #9ca3af;
        --border-color:  #1f2937;
        --accent:        #4f46e5;
        --accent-hover:  #6366f1;
    }

    html, body, [data-testid="stAppViewContainer"] {
        background: var(--bg-primary) !important;
        color: var(--text-primary) !important;
        font-family: 'Inter', sans-serif !important;
    }
    [data-testid="stHeader"] { background: transparent !important; }
    [data-testid="stSidebar"] {
        background: var(--bg-secondary) !important;
        border-right: 1px solid var(--border-color) !important;
    }
    [data-testid="stSidebar"] * { color: var(--text-secondary) !important; }

    ::-webkit-scrollbar { width: 3px; }
    ::-webkit-scrollbar-track { background: transparent; }
    ::-webkit-scrollbar-thumb { background: #1e2030; border-radius: 3px; }

    #MainMenu, footer, [data-testid="stDecoration"] { display: none !important; }

    div[data-testid="stButton"] > button {
        font-family: 'Inter', sans-serif !important;
        transition: all 0.15s ease !important;
    }
    div[data-testid="stButton"] > button:focus { box-shadow: none !important; }

    div[data-testid="stButton"] > button[kind="primary"] {
        background: #3b5bdb !important;
        border: none !important;
        color: #fff !important;
        border-radius: 8px !important;
        font-weight: 500 !important;
        font-size: 0.82rem !important;
        font-family: 'Inter', sans-serif !important;
        padding: 8px 16px !important;
    }
    div[data-testid="stButton"] > button[kind="primary"]:hover {
        background: #3451c7 !important;
    }
    div[data-testid="stButton"] > button[kind="secondary"] {
        background: transparent !important;
        border: 1px solid #1c1e2a !important;
        color: var(--text-secondary) !important;
        border-radius: 8px !important;
        font-size: 0.82rem !important;
        font-weight: 400 !important;
        padding: 6px 14px !important;
        min-height: unset !important;
    }
    div[data-testid="stButton"] > button[kind="secondary"]:hover {
        border-color: #2e3248 !important;
        color: var(--text-primary) !important;
    }

    /* Download button — Zeus style */
    [data-testid="stDownloadButton"] > button {
        background: transparent !important;
        border: 1px solid #3b5bdb !important;
        color: #748ffc !important;
        border-radius: 8px !important;
        font-size: 0.82rem !important;
        font-family: 'Inter', sans-serif !important;
        font-weight: 500 !important;
        padding: 8px 16px !important;
    }
    [data-testid="stDownloadButton"] > button:hover {
        background: rgba(59,91,219,0.08) !important;
    }

    [data-testid="stFileUploader"] {
        background: var(--bg-secondary) !important;
        border: 1px dashed var(--border-color) !important;
        border-radius: 8px !important;
        padding: 8px !important;
    }
    [data-testid="stFileUploader"]:hover { border-color: var(--accent) !important; }

    [data-testid="stMetric"] {
        background: var(--bg-secondary) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 8px !important;
        padding: 12px !important;
    }
    [data-testid="stAlert"] {
        background: var(--bg-secondary) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 8px !important;
        color: var(--text-secondary) !important;
    }

    /* Select boxes */
    div[data-baseweb="select"] > div {
        background: var(--bg-secondary) !important;
        border-color: var(--border-color) !important;
        border-radius: 8px !important;
        color: var(--text-primary) !important;
        font-family: 'Inter', sans-serif !important;
        font-size: 0.82rem !important;
    }
    div[data-baseweb="select"] > div:focus-within {
        border-color: var(--accent) !important;
        box-shadow: 0 0 0 2px rgba(79,70,229,0.12) !important;
    }

    /* Slider */
    .stSlider > div > div { background: var(--border-color) !important; }
    .stSlider [data-baseweb="slider"] div[role="slider"] {
        background: #3b5bdb !important;
        border-color: #3b5bdb !important;
    }

    /* Expander */
    [data-testid="stExpander"] {
        background: var(--card-bg) !important;
        border: 1px solid var(--border-color) !important;
        border-radius: 8px !important;
    }
    [data-testid="stExpander"] summary {
        color: var(--text-secondary) !important;
        font-size: 0.82rem !important;
    }

    /* Caption */
    [data-testid="stCaptionContainer"] {
        color: var(--text-muted) !important;
        font-size: 0.75rem !important;
    }

    /* Progress bar */
    [data-testid="stProgressBar"] > div > div {
        background: #3b5bdb !important;
    }

    /* ── shared component styles (Zeus Doc Chat) ── */
    .zdc-section-label {
        font-size: 10px;
        font-weight: 600;
        text-transform: uppercase;
        letter-spacing: 0.1em;
        color: var(--text-secondary);
        margin-bottom: 10px;
        margin-top: 20px;
    }
    .zdc-section-label:first-child { margin-top: 0; }

    .zdc-status-ok {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        background: rgba(34,197,94,0.06);
        border: 1px solid rgba(34,197,94,0.18);
        color: var(--text-primary);
        font-size: 0.72rem;
        font-weight: 500;
        padding: 5px 12px;
        border-radius: 100px;
        margin-top: 8px;
    }
    .zdc-status-ok::before {
        content: '';
        width: 6px; height: 6px;
        background: #4ade80;
        border-radius: 50%;
        display: inline-block;
    }

    .zdc-page-header {
        display: flex;
        align-items: center;
        gap: 14px;
        padding: 20px 0 18px;
        border-bottom: 1px solid #1c1e2a;
        margin-bottom: 20px;
    }
    .zdc-page-icon {
        width: 38px; height: 38px;
        background: #13141d;
        border: 1px solid var(--border-color);
        border-radius: 8px;
        display: flex; align-items: center; justify-content: center;
        color: #5c7cfa;
        font-size: 1rem;
        font-weight: 600;
        flex-shrink: 0;
    }
    .zdc-page-title {
        font-size: 1.05rem !important;
        font-weight: 600 !important;
        color: var(--text-primary) !important;
        margin: 0 !important;
        letter-spacing: -0.01em;
    }
    .zdc-rag-badge {
        display: inline-flex;
        align-items: center;
        gap: 4px;
        background: rgba(59,91,219,0.08);
        border: 1px solid rgba(59,91,219,0.2);
        color: var(--text-primary);
        font-size: 0.68rem;
        font-weight: 500;
        padding: 2px 9px;
        border-radius: 100px;
        margin-left: 10px;
        vertical-align: middle;
    }

    .zdc-empty {
        text-align: center;
        padding: 60px 20px;
    }
    .zdc-empty-icon {
        width: 44px; height: 44px;
        background: #13141d;
        border: 1px solid var(--border-color);
        border-radius: 10px;
        display: flex; align-items: center; justify-content: center;
        margin: 0 auto 14px;
        color: var(--text-muted);
        font-size: 1.1rem;
        font-weight: 600;
    }
    .zdc-empty-title { font-size: 0.9rem; font-weight: 500; color: var(--text-primary); margin-bottom: 4px; }
    .zdc-empty-sub   { font-size: 0.78rem; color: var(--text-secondary); }

    /* Result rows */
    .zdc-result-row {
        display: flex;
        align-items: flex-start;
        gap: 10px;
        background: var(--card-bg);
        border: 1px solid var(--border-color);
        border-radius: 8px;
        padding: 10px 14px;
        margin-bottom: 6px;
        font-size: 0.8rem;
        color: var(--text-secondary);
    }
    .zdc-result-row.success { border-left: 3px solid #4ade80; }
    .zdc-result-row.skipped { border-left: 3px solid #6b7280; }
    .zdc-result-row.failed  { border-left: 3px solid #f87171; color: #f87171; }
    .zdc-result-source { font-weight: 500; color: var(--text-primary); }
    .zdc-result-note   { color: var(--text-muted); font-size: 0.75rem; margin-top: 2px; }

    /* Stat cards */
    .zdc-stat-row { display: flex; gap: 12px; margin: 16px 0; }
    .zdc-stat-card {
        flex: 1;
        background: var(--card-bg);
        border: 1px solid var(--border-color);
        border-radius: 10px;
        padding: 14px 16px;
        text-align: center;
    }
    .zdc-stat-num   { font-size: 1.8rem; font-weight: 600; line-height: 1; }
    .zdc-stat-label { font-size: 0.68rem; text-transform: uppercase; letter-spacing: 0.1em; color: var(--text-muted); margin-top: 4px; }
    .zdc-stat-card.green .zdc-stat-num { color: #4ade80; }
    .zdc-stat-card.gray  .zdc-stat-num { color: #6b7280; }
    .zdc-stat-card.red   .zdc-stat-num { color: #f87171; }

    /* Format badges */
    .zdc-badge {
        display: inline-block;
        font-size: 10px;
        font-weight: 500;
        padding: 2px 7px;
        border-radius: 4px;
        background: #13141d;
        border: 1px solid var(--border-color);
        color: var(--text-secondary);
        margin: 2px;
    }
    .zdc-badge.warn { border-color: rgba(251,191,36,0.25); color: #fbbf24; }

    .zdc-format-group { margin-bottom: 10px; }
    .zdc-format-group-title {
        font-size: 10px;
        text-transform: uppercase;
        letter-spacing: 0.1em;
        color: var(--text-muted);
        margin-bottom: 5px;
        font-weight: 600;
    }

    /* Quality box */
    .zdc-quality-box {
        background: var(--card-bg);
        border: 1px solid var(--border-color);
        border-radius: 8px;
        padding: 14px 16px;
        margin: 12px 0;
    }
    .zdc-quality-label {
        font-size: 0.75rem;
        font-weight: 500;
        color: var(--text-secondary);
        margin-bottom: 8px;
    }
    </style>
    """, unsafe_allow_html=True)


# ============================================================
# STREAMLIT UI — ZEUS DOC CHAT STYLE
# ============================================================

inject_global_css()

FORMAT_MAP = {
    "🖼️ Image":    ["webp", "png", "jpg", "jpeg", "bmp", "gif", "tiff"],
    "🎵 Audio":    ["mp3", "wav", "ogg", "flac", "aac"],
    "📄 Document": ["txt", "docx", "pdf", "html", "md", "rtf", "odt", "csv", "xlsx"],
}

QUALITY_FORMATS = {"jpg", "jpeg", "webp", "mp3", "ogg", "aac"}
IS_BITRATE      = {"mp3", "ogg", "aac"}

# ── SIDEBAR ────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="zdc-section-label">Files</div>', unsafe_allow_html=True)

    uploaded_files = st.file_uploader(
        "Upload your files",
        accept_multiple_files=True,
        label_visibility="collapsed"
    )

    if uploaded_files:
        st.markdown(
            f'<div class="zdc-status-ok">{len(uploaded_files)} file(s) ready</div>',
            unsafe_allow_html=True
        )

    st.markdown('<div class="zdc-section-label">Target Format</div>', unsafe_allow_html=True)

    file_type_label = st.selectbox(
        "File Type",
        list(FORMAT_MAP.keys()),
        label_visibility="collapsed"
    )
    target_format = st.selectbox(
        "Format",
        FORMAT_MAP[file_type_label],
        label_visibility="collapsed"
    )

    quality = None
    if target_format in QUALITY_FORMATS:
        st.markdown('<div class="zdc-section-label">Quality</div>', unsafe_allow_html=True)
        if target_format in IS_BITRATE:
            quality = st.slider("Bitrate (kbps)", 64, 320,
                                int(DEFAULT_QUALITY.get(target_format, 192)), step=32)
            st.caption(f"Audio bitrate: **{quality} kbps**")
        else:
            quality = st.slider("Image Quality", 1, 100,
                                int(DEFAULT_QUALITY.get(target_format, 85)))
            st.caption(f"Quality: **{quality}/100**")

    st.markdown('<div class="zdc-section-label">Actions</div>', unsafe_allow_html=True)
    convert_btn = st.button("Convert All Files", use_container_width=True, type="primary")


# ── MAIN AREA ──────────────────────────────────────────────────────────────────

st.markdown("""
<div class="zdc-page-header">
  <div class="zdc-page-icon">⚡</div>
  <div>
    <div class="zdc-page-title">ZEUS Bulk Format Converter
      <span class="zdc-rag-badge">Images · Audio · Docs</span>
    </div>
  </div>
</div>
""", unsafe_allow_html=True)

# Supported formats reference
with st.expander("Supported Input Formats"):
    st.markdown("""
    <div class="zdc-format-group">
      <div class="zdc-format-group-title">🖼️ Images</div>
      <span class="zdc-badge">.png</span><span class="zdc-badge">.jpg</span>
      <span class="zdc-badge">.jpeg</span><span class="zdc-badge">.bmp</span>
      <span class="zdc-badge">.gif</span><span class="zdc-badge">.tiff</span>
      <span class="zdc-badge">.webp</span><span class="zdc-badge">.ico</span>
      <span class="zdc-badge">.ppm</span>
    </div>
    <div class="zdc-format-group">
      <div class="zdc-format-group-title">🎵 Audio</div>
      <span class="zdc-badge">.mp3</span><span class="zdc-badge">.wav</span>
      <span class="zdc-badge">.ogg</span><span class="zdc-badge">.flac</span>
      <span class="zdc-badge">.aac</span><span class="zdc-badge">.wma</span>
      <span class="zdc-badge">.m4a</span>
    </div>
    <div class="zdc-format-group">
      <div class="zdc-format-group-title">📄 Documents — Word</div>
      <span class="zdc-badge">.docx</span>
      <span class="zdc-badge warn">.doc ⚠ LibreOffice</span>
      <span class="zdc-badge">.dotx</span><span class="zdc-badge">.dotm</span>
      <span class="zdc-badge">.docm</span>
    </div>
    <div class="zdc-format-group">
      <div class="zdc-format-group-title">📄 Documents — Other</div>
      <span class="zdc-badge">.pdf</span><span class="zdc-badge">.rtf</span>
      <span class="zdc-badge">.odt</span><span class="zdc-badge">.txt</span>
      <span class="zdc-badge">.md</span><span class="zdc-badge">.html</span>
      <span class="zdc-badge">.htm</span><span class="zdc-badge">.csv</span>
      <span class="zdc-badge">.xlsx</span>
    </div>
    """, unsafe_allow_html=True)

# Empty state
if not uploaded_files:
    st.markdown("""
    <div class="zdc-empty">
      <div class="zdc-empty-icon">[ ]</div>
      <div class="zdc-empty-title">No files uploaded</div>
      <div class="zdc-empty-sub">Upload files and choose a target format using the sidebar, then hit Convert</div>
    </div>
    """, unsafe_allow_html=True)
else:
    # File list preview
    st.markdown('<div class="zdc-section-label">Uploaded Files</div>', unsafe_allow_html=True)
    names_preview = [f.name for f in uploaded_files[:6]]
    extra = len(uploaded_files) - 6
    for name in names_preview:
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else "?"
        st.markdown(
            f'<div class="zdc-result-row">'
            f'<span class="zdc-badge">.{ext}</span>'
            f'<span class="zdc-result-source">{name}</span>'
            f'</div>',
            unsafe_allow_html=True
        )
    if extra > 0:
        st.caption(f"…and {extra} more file(s)")

# ── CONVERSION ─────────────────────────────────────────────────────────────────
if convert_btn:
    if not uploaded_files:
        st.warning("Please upload at least one file first.")
        st.stop()

    results = []
    converted_files = {}
    progress = st.progress(0, text="Starting…")
    total = len(uploaded_files)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp     = Path(tmpdir)
        out_dir = tmp / "out"
        out_dir.mkdir()

        for i, uf in enumerate(uploaded_files):
            progress.progress(i / total, text=f"Converting {uf.name}…")
            in_path = tmp / uf.name
            in_path.write_bytes(uf.read())

            ext       = in_path.suffix.lower()
            file_type = detect_file_type(ext)

            if file_type is None:
                results.append(ConversionResult(uf.name, "", "skipped", f"Unrecognised format ({ext})"))
                continue
            if ext == f".{target_format}":
                results.append(ConversionResult(uf.name, "", "skipped", "Already in target format"))
                continue
            if target_format not in get_valid_outputs(file_type):
                results.append(ConversionResult(uf.name, "", "skipped", f"Cannot convert {file_type} → .{target_format}"))
                continue

            out_path = out_dir / f"{in_path.stem}.{target_format}"
            c = 1
            while out_path.exists():
                out_path = out_dir / f"{in_path.stem}_{c}.{target_format}"
                c += 1

            try:
                if file_type == "image":
                    convert_image(in_path, out_path, quality)
                elif file_type == "audio":
                    convert_audio(in_path, out_path, quality)
                elif file_type == "doc":
                    convert_doc(in_path, out_path)
                converted_files[out_path.name] = out_path.read_bytes()
                results.append(ConversionResult(uf.name, out_path.name, "success"))
            except Exception as e:
                results.append(ConversionResult(uf.name, "", "failed", str(e)))

        progress.progress(1.0, text="Done!")

    # ── Summary stats ──────────────────────────────────────────────────────────
    success = sum(1 for r in results if r.status == "success")
    skipped = sum(1 for r in results if r.status == "skipped")
    failed  = sum(1 for r in results if r.status == "failed")

    st.markdown(f"""
    <div class="zdc-stat-row">
      <div class="zdc-stat-card green">
        <div class="zdc-stat-num">{success}</div>
        <div class="zdc-stat-label">Converted</div>
      </div>
      <div class="zdc-stat-card gray">
        <div class="zdc-stat-num">{skipped}</div>
        <div class="zdc-stat-label">Skipped</div>
      </div>
      <div class="zdc-stat-card red">
        <div class="zdc-stat-num">{failed}</div>
        <div class="zdc-stat-label">Failed</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # ── Per-file results ───────────────────────────────────────────────────────
    st.markdown('<div class="zdc-section-label">Results</div>', unsafe_allow_html=True)
    for r in results:
        icon = {"success": "✅", "skipped": "⏭️", "failed": "❌"}.get(r.status, "•")
        css  = r.status
        note = r.error if r.error else (f"→ {r.target}" if r.target else "")
        st.markdown(
            f'<div class="zdc-result-row {css}">'
            f'<span style="font-size:0.9rem">{icon}</span>'
            f'<div>'
            f'<div class="zdc-result-source">{r.source}</div>'
            f'<div class="zdc-result-note">{note}</div>'
            f'</div>'
            f'</div>',
            unsafe_allow_html=True
        )

    # ── Downloads ──────────────────────────────────────────────────────────────
    if converted_files:
        st.markdown("<br>", unsafe_allow_html=True)
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for fname, data in converted_files.items():
                zf.writestr(fname, data)
        zip_buf.seek(0)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        st.download_button(
            label=f"⬇ Download {len(converted_files)} converted file(s) as ZIP",
            data=zip_buf,
            file_name=f"converted_{ts}.zip",
            mime="application/zip",
            use_container_width=True,
        )

    if results:
        ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
        log = [f"Conversion Log — {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n" + "=" * 60]
        for r in results:
            log.append(f"\n[{r.status.upper()}]\n  Source: {r.source}")
            if r.target: log.append(f"  Output: {r.target}")
            if r.error:  log.append(f"  Note:   {r.error}")
        st.download_button(
            label="📋 Download conversion log",
            data="\n".join(log),
            file_name=f"conversion_log_{ts}.txt",
            mime="text/plain",
        )
